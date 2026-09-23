from __future__ import annotations

import ast
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
ANNOTATED = ROOT / "assets" / "ui-annotated"
OUT = ROOT / "GRAV-Store-Employee-Quick-Help-Guide.docx"
ANNOTATION_SOURCE = ROOT / "build_ui_annotations.py"

BLACK = "000000"
NAVY = "17365D"
PALE_BLUE = "EAF1FB"
PALE_GRAY = "F5F6F8"
GRID = "D9D9D9"
MUTED = RGBColor(82, 91, 104)


ROUTES = {
    "01-overview": "Overview",
    "02-request-desk": "Requests > Request desk",
    "25-request-detail-open": "Requests > Request desk > View request",
    "29-request-required-items": "Request detail > Required Items",
    "32-issue-stock-form": "Request detail > Issue Stock",
    "33-return-stock-form": "Request detail > Return Stock",
    "03-purchase-orders": "Purchase > Purchase orders",
    "27-new-purchase-order": "Purchase > Purchase orders > New purchase order",
    "05-goods-receipts": "Receive > Goods receipts",
    "04-service-orders": "Purchase > Service orders",
    "06-reservations-picking": "Inventory > Reservations and picking",
    "07-issues-returns": "Inventory > Issues and returns",
    "08-stock-counts": "Inventory > Stock counts",
    "09-stock-movements": "Inventory > Stock movements",
    "10-stock-exceptions": "Inventory > Stock exceptions",
    "11-lot-labels": "Inventory > Lot labels and barcodes",
    "12-materials": "Masters > Materials",
    "28-new-material": "Masters > Materials > Add item",
    "13-finished-products": "Masters > Finished products and BOM",
    "15-suppliers": "Masters > Suppliers",
    "16-supplier-offers": "Masters > Supplier offers",
    "35-new-supplier-offer": "Masters > Supplier offers > New offer",
    "17-units": "Masters > Units of measure",
    "18-warehouses": "Masters > Warehouses",
    "38-new-warehouse-form": "Masters > Warehouses > Add warehouse",
    "19-request-report": "Reports > Request report",
    "20-movement-report": "Reports > Movement report",
    "21-purchase-exceptions": "Reports > Purchase exceptions",
    "22-inventory-valuation": "Reports > Inventory valuation",
    "23-settings": "Settings",
    "24-legacy-purchase-forms": "Legacy and production > Purchase forms legacy",
}

WHEN = {
    "01-overview": "At the start of a shift or whenever you need to decide what requires attention next.",
    "02-request-desk": "An approved material or service need must be reviewed and routed to stock, purchase or service fulfilment.",
    "25-request-detail-open": "Before reserving, issuing, buying or ordering a service against a request.",
    "29-request-required-items": "The request contains product demand that must be translated into exact material and variant requirements.",
    "32-issue-stock-form": "Approved material is physically leaving Store for a known request or work context.",
    "33-return-stock-form": "Unused material previously issued is physically returning to Store.",
    "03-purchase-orders": "You need to find, monitor, issue, receive against or investigate a physical-goods order.",
    "27-new-purchase-order": "Approved demand is ready to become a supplier-facing order for physical goods.",
    "05-goods-receipts": "Goods have arrived against a purchase order and require quantity and condition recording.",
    "04-service-orders": "Non-stock work is issued to a supplier and must be followed through completion and requester acceptance.",
    "06-reservations-picking": "Approved demand should be promised, picked and then issued from usable stock.",
    "07-issues-returns": "You need the operational workspace for a stock-out or an internal stock return.",
    "08-stock-counts": "Recorded stock must be checked against a physical location and differences controlled.",
    "09-stock-movements": "You need to explain how an item's recorded quantity changed.",
    "10-stock-exceptions": "The system reports a negative balance, placement mismatch or broken stock evidence chain.",
    "11-lot-labels": "A stored quantity needs a traceable printed label or barcode.",
    "12-materials": "You need to find or review a stock-tracked input item without changing its quantity.",
    "28-new-material": "A genuinely new reusable material identity is required and no duplicate exists.",
    "13-finished-products": "You need production-owned finished-product identity or BOM context.",
    "15-suppliers": "You need to find or maintain the legal and operational identity of a supplier.",
    "16-supplier-offers": "A sourcing or costing decision needs dated commercial evidence from suppliers.",
    "35-new-supplier-offer": "A new quotation must be recorded against the exact supplier and item or service.",
    "17-units": "A controlled unit or defensible conversion is missing from master data.",
    "18-warehouses": "You need to review the company's physical facilities and operational locations.",
    "38-new-warehouse-form": "A real physical facility is ready to be represented in Store.",
    "19-request-report": "You need to analyse demand, ageing or incomplete request fulfilment over a defined scope.",
    "20-movement-report": "You need period-based movement activity or source-document reconciliation.",
    "21-purchase-exceptions": "Orders, receipts, bills or service records are overdue, mismatched or blocked.",
    "22-inventory-valuation": "Management needs inventory value, evidence coverage and missing-cost visibility.",
    "23-settings": "An administrator must change future Store document identity or operating defaults.",
    "24-legacy-purchase-forms": "Historical purchase-form evidence must be read; never use it to begin new work.",
}

OUTCOME = {
    "01-overview": "You enter the correct filtered queue and know which work owns the next action.",
    "02-request-desk": "The correct request is open and its fulfilment path is visible.",
    "25-request-detail-open": "Every action is based on the request's current approval, identity and required date.",
    "29-request-required-items": "Requested descriptions point to the correct material, variant, quantity and unit.",
    "32-issue-stock-form": "One traceable stock-out is recorded against the correct request.",
    "33-return-stock-form": "Returned quantity reaches the correct usable or non-usable location with a reason.",
    "03-purchase-orders": "The order and its next operational action are identified.",
    "27-new-purchase-order": "A reviewable draft exists; issuing remains a separate controlled decision.",
    "05-goods-receipts": "Accepted, quarantined and rejected quantities remain separate and traceable.",
    "04-service-orders": "Accepted service evidence is ready for Accounting without creating stock or a goods receipt.",
    "06-reservations-picking": "Usable stock is promised and picked without hiding shortages or issuing twice.",
    "07-issues-returns": "Direction, item, variant, quantity, unit and reason match the physical event.",
    "08-stock-counts": "Observations and approved corrections remain auditable.",
    "09-stock-movements": "The source document and movement chain explain the quantity change.",
    "10-stock-exceptions": "The source condition is corrected without rewriting posted history.",
    "11-lot-labels": "Each label identifies a real quantity without changing inventory.",
    "12-materials": "The existing material is found or a genuine master-data gap is confirmed.",
    "28-new-material": "A stable material identity exists with no invented opening balance.",
    "13-finished-products": "The correct product and BOM are understood without duplicating production master data.",
    "15-suppliers": "Purchasing records refer to one controlled supplier identity.",
    "16-supplier-offers": "The commercial decision is supported by comparable, current evidence.",
    "35-new-supplier-offer": "A dated and attributable quotation can support sourcing or costing.",
    "17-units": "Quantities retain a valid and explainable measurement basis.",
    "18-warehouses": "Transactions can point to the place where stock physically exists.",
    "38-new-warehouse-form": "The facility and its standard locations are ready for controlled use.",
    "19-request-report": "The result can be explained by its visible filters and source requests.",
    "20-movement-report": "Movement activity is reconciled without combining unlike units.",
    "21-purchase-exceptions": "The actual blocker has an owner and a verifiable resolution path.",
    "22-inventory-valuation": "Value is read together with method, coverage and limitations.",
    "23-settings": "The change is verified, documented and effective from a known point.",
    "24-legacy-purchase-forms": "Historical evidence remains readable while new work begins in Request desk.",
}

BLOCKED = {
    "01-overview": "If a card and its register disagree, refresh once and use the register as the working record.",
    "02-request-desk": "If the request is missing, verify company, date, requester and stage before escalating.",
    "25-request-detail-open": "If actions are hidden, check status and role; do not use another person's sign-in.",
    "29-request-required-items": "If calculation remains incomplete, refresh the request and do not act on partial quantities.",
    "32-issue-stock-form": "If stock or a conversion is missing, stop; correct availability or master data first.",
    "33-return-stock-form": "If the original issue cannot be identified, do not credit an arbitrary item.",
    "03-purchase-orders": "If an issued order is wrong, use the supported amendment or cancellation path; do not overwrite history.",
    "27-new-purchase-order": "If demand, supplier, budget or unit evidence is incomplete, keep the order in Draft.",
    "05-goods-receipts": "If the PO line, unit or location is unclear, quarantine or pause instead of guessing.",
    "04-service-orders": "If completion cannot be accepted, the requester should ask for rework with a reason.",
    "06-reservations-picking": "If availability is short, leave the shortage visible; never make stock negative.",
    "07-issues-returns": "If the physical event and selected direction differ, cancel before submitting.",
    "08-stock-counts": "If warehouse structure is missing, ask an administrator to configure it before counting.",
    "09-stock-movements": "If the chain is suspicious, open the source record; never edit a posted movement.",
    "10-stock-exceptions": "If evidence is insufficient, start a controlled physical count rather than clearing the warning.",
    "11-lot-labels": "Reprinting does not create stock; if quantity is wrong, correct the underlying transaction.",
    "12-materials": "If an item looks similar, compare code, variant and base unit before creating another.",
    "28-new-material": "If category, base unit or budget classification is uncertain, save nothing and ask the master-data owner.",
    "13-finished-products": "Engineering changes belong in Production, not in the Store material master.",
    "15-suppliers": "If a possible duplicate exists, resolve it before adding a new supplier.",
    "16-supplier-offers": "If validity or unit basis is missing, the offer is not comparable; obtain corrected evidence.",
    "35-new-supplier-offer": "If supplier, currency, unit or validity is unknown, do not invent it.",
    "17-units": "If no valid conversion exists, add one only with defensible business evidence.",
    "18-warehouses": "If a warehouse appears only under legacy records, ask the administrator before new use.",
    "38-new-warehouse-form": "Do not create a warehouse as a substitute for an internal bin or temporary stock state.",
    "19-request-report": "If totals look unexpected, clear filters and read the population-scope sentence.",
    "20-movement-report": "Do not treat filtered net movement as current on-hand stock.",
    "21-purchase-exceptions": "Correct the source record, then refresh; do not merely dismiss the symptom.",
    "22-inventory-valuation": "Missing cost evidence is unknown, not zero; resolve receipt or landed-cost evidence.",
    "23-settings": "If downstream document impact is unclear, do not save until Finance and Store owners agree.",
    "24-legacy-purchase-forms": "For new demand, leave this screen and use Request desk.",
}


def load_specs():
    tree = ast.parse(ANNOTATION_SOURCE.read_text())
    for node in tree.body:
        if isinstance(node, ast.Assign) and any(isinstance(t, ast.Name) and t.id == "SPECS" for t in node.targets):
            return ast.literal_eval(node.value)
    raise RuntimeError("SPECS not found")


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def borders(table):
    tbl_pr = table._tbl.tblPr
    item = tbl_pr.find(qn("w:tblBorders"))
    if item is None:
        item = OxmlElement("w:tblBorders")
        tbl_pr.append(item)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = item.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            item.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:color"), GRID)


def margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc.get_or_add_tcPr()
    tc_mar = tc.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        el = tc_mar.find(qn(f"w:{name}"))
        if el is None:
            el = OxmlElement(f"w:{name}")
            tc_mar.append(el)
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, end))


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.55)
section.bottom_margin = Inches(0.55)
section.left_margin = Inches(0.65)
section.right_margin = Inches(0.65)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.08

for style_name, size in (("Title", 28), ("Subtitle", 13), ("Heading 1", 21), ("Heading 2", 16), ("Heading 3", 12)):
    style = styles[style_name]
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.bold = style_name != "Subtitle"
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(6)
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.keep_with_next = True

footer = section.footer.paragraphs[0]
footer.add_run("GRAV Store Employee Quick Help  |  ").font.size = Pt(8)
add_page_number(footer)


def p(text="", bold_lead=None, style=None):
    para = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        para.add_run(bold_lead).bold = True
        para.add_run(text[len(bold_lead):])
    else:
        para.add_run(text)
    return para


def bullets(items, numbered=False):
    style = "List Number" if numbered else "List Bullet"
    for item in items:
        para = doc.add_paragraph(style=style)
        para.paragraph_format.space_after = Pt(3)
        para.add_run(item)


def table(headers, rows, widths=None, font_size=9.2):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.autofit = False
    borders(tbl)
    for i, head in enumerate(headers):
        cell = tbl.rows[0].cells[i]
        shade(cell, NAVY)
        margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            cell.width = Inches(widths[i])
        run = cell.paragraphs[0].add_run(head)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(font_size)
    for ridx, row in enumerate(rows):
        cells = tbl.add_row().cells
        if ridx % 2:
            for cell in cells:
                shade(cell, PALE_GRAY)
        for i, value in enumerate(row):
            cell = cells[i]
            margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                cell.width = Inches(widths[i])
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            run = para.add_run(str(value))
            run.font.size = Pt(font_size)
        tbl.rows[-1].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    return tbl


def new_page():
    doc.add_page_break()


# Cover
doc.add_paragraph("GRAV Store Employee Quick Help", style="Title")
doc.add_paragraph("Find an answer and complete one task without reading the full handbook", style="Subtitle")
p("This guide is for employees who need to complete Store and Purchase work quickly and correctly. It is not designed to be read from beginning to end. Start with your question, role or current screen, then open only the task card you need.")
doc.add_picture(str(ANNOTATED / "01-overview-annotated.png"), width=Inches(7.15))
p("Daily rule", style="Heading 2")
p("Search first, verify the exact business record, perform one controlled action, then confirm the resulting status or movement.")
p("Detailed policy and unusual cases remain in the GRAV Store and Purchase Complete User Guide.")

new_page()
doc.add_paragraph("Find an answer in 30 seconds", style="Heading 1")
p("Use whichever starting point matches what you know. Word search and the Navigation pane work throughout this guide.")
table(
    ["What you know", "What to do", "Example"],
    [
        ("Your goal", "Search the verb or object, then open the matching task card.", "receive, return, supplier, count, valuation"),
        ("Your role", "Open My work by role and follow only your normal queue.", "requester, buyer, receiver, inventory controller"),
        ("Your current screen", "Search the visible page title exactly as it appears.", "Goods receipts, Stock exceptions, Materials"),
        ("An error message", "Search a distinctive phrase from the message and use Problem to action.", "unit conversion, insufficient stock, staff record"),
        ("A document number", "Open the owning register and search the exact number.", "request, PO, receipt, service order"),
    ],
    [1.35, 3.25, 2.45],
)
doc.add_paragraph("Do not browse module names at random", style="Heading 2")
bullets([
    "A new need always begins in Request desk, not in a legacy register.",
    "Physical goods use Purchase orders and Goods receipts.",
    "Non-stock work uses Service orders and requester acceptance; it creates no stock.",
    "Quantity changes belong to receiving, issue, return, transfer, count or correction workflows - not master forms.",
    "Budget mapping and Accounting actuals are connected controls, not fields employees should invent while handling stock.",
])

new_page()
doc.add_paragraph("Goal index", style="Heading 1")
table(
    ["I want to", "Open", "Search words"],
    [
        ("See what needs attention", "Card 01 Overview", "shift, urgent, queue"),
        ("Find an approved need", "Cards 02 to 04", "request, required items, classify"),
        ("Issue or return material", "Cards 05, 06 and 12", "issue, return, credit stock"),
        ("Buy physical goods", "Cards 07 and 08", "purchase order, supplier, PO"),
        ("Receive and inspect goods", "Card 09", "receipt, GRN, quarantine, reject"),
        ("Order and accept a service", "Card 10", "service order, completion, rework"),
        ("Reserve or pick stock", "Card 11", "reservation, available, picked"),
        ("Count or investigate stock", "Cards 13 to 15", "count, movement, exception"),
        ("Print a traceable label", "Card 16", "lot, barcode, label"),
        ("Find or create master data", "Cards 17 to 25", "material, product, supplier, offer, unit, warehouse"),
        ("Read operational reports", "Cards 26 to 29", "request report, movement, exception, value"),
        ("Change Store settings", "Card 30", "address, PDF, settings"),
        ("Find historical purchase forms", "Card 31", "legacy, old purchase form, history"),
    ],
    [2.1, 2.25, 2.7],
)
p("Card numbers in this guide are stable lookup identifiers. Search CARD plus the two-digit number, or search the page title shown in the Open column.")

new_page()
doc.add_paragraph("My work by role", style="Heading 1")
table(
    ["Role", "Start here", "Normal work", "Stop and escalate when"],
    [
        ("Requester", "Request desk", "Raise and follow a need; provide purpose and required date; accept completed services.", "The request identity, quantity or service result is wrong."),
        ("Approver", "Request detail", "Confirm business need, timing and scope; approve or refuse with a reason.", "Evidence, authority or budget result is unclear."),
        ("Buyer", "Request desk / Purchase orders", "Source approved demand; compare offers; prepare and issue supplier documents.", "Supplier, unit, tax, price or budget evidence is incomplete."),
        ("Store receiver", "Overview / Goods receipts", "Match arrival to PO; record quantity and condition; quarantine or put away.", "PO line, unit, location or condition cannot be verified."),
        ("Inventory controller", "Reservations / Counts / Exceptions", "Reserve, pick, issue, return, count and resolve evidence breaks.", "A correction would require editing posted history or making stock negative."),
        ("Manager", "Overview / Reports", "Prioritise queues, review exceptions and confirm handoffs have owners.", "Operational and financial records disagree without evidence."),
        ("Finance reader", "Valuation / source documents", "Read management value with coverage; match commercial and accounting evidence.", "Missing evidence is being treated as zero cost."),
        ("Administrator", "Masters / Settings", "Maintain controlled identities, units, locations and future defaults.", "A change could rewrite or misclassify historical work."),
    ],
    [1.05, 1.4, 2.85, 1.75],
    8.6,
)

new_page()
doc.add_paragraph("Start and finish the shift", style="Heading 1")
doc.add_paragraph("Start of shift", style="Heading 2")
bullets([
    "Open Overview and read Needs attention before routine registers.",
    "Open the highest-priority filtered queue from its card.",
    "Check receipts awaiting inspection, put-away pending, reservation shortages and stock exceptions.",
    "Read yesterday's unresolved notes and confirm ownership before creating new work.",
])
doc.add_paragraph("Before every stock-changing action", style="Heading 2")
bullets([
    "Verify document number, item, variant, quantity, unit and physical location.",
    "Confirm the screen's action matches the real event: receive, reserve, pick, issue, return, transfer, count or correct.",
    "Submit once. If the response is uncertain, refresh and inspect history before trying again.",
])
doc.add_paragraph("End of shift", style="Heading 2")
bullets([
    "No accepted receipt remains without a known location unless it is deliberately quarantined.",
    "Picked material is either issued, returned to usable stock or left with a named owner.",
    "Open counts and unresolved exceptions have a reason and next owner.",
    "The physical handover agrees with the system's visible queues.",
])


specs = load_specs()
for idx, (stem, title, steps) in enumerate(specs, 1):
    new_page()
    doc.add_paragraph(f"CARD {idx:02d}  {title}", style="Heading 1")
    p(f"Open: {ROUTES[stem]}", bold_lead="Open:")
    p(f"Use this when: {WHEN[stem]}", bold_lead="Use this when:")
    image = ANNOTATED / f"{stem}-annotated.png"
    if not image.exists():
        raise FileNotFoundError(image)
    pic = doc.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.add_run().add_picture(str(image), width=Inches(7.0))
    step_p = doc.add_paragraph()
    step_p.paragraph_format.space_before = Pt(2)
    step_p.paragraph_format.space_after = Pt(4)
    step_p.add_run("Steps: ").bold = True
    step_p.add_run("  ".join(f"{n}. {copy}" for n, (copy, _anchor) in enumerate(steps, 1)))
    p(f"Expected result: {OUTCOME[stem]}", bold_lead="Expected result:")
    p(f"If blocked: {BLOCKED[stem]}", bold_lead="If blocked:")


new_page()
doc.add_paragraph("Problem to action", style="Heading 1")
p("Search the exact words shown on screen. Do not repeatedly retry a refusal whose cause requires data, permission or another person's action.")
table(
    ["Problem or message", "What it usually means", "What to do next"],
    [
        ("Nothing appears in a register", "Filters, date scope, company scope or the wrong register may hide it.", "Clear filters; search exact number; confirm status and owning register; refresh once."),
        ("Action is hidden", "Your role, capability or the record's status does not permit it.", "Confirm role and workflow state. Ask the owner; never use another login."),
        ("Staff record not linked", "The sign-in can read some registers but cannot be named in an audited mutation.", "Ask an administrator to link the account to the correct employee record."),
        ("Insufficient stock", "Available usable stock is below the requested quantity.", "Check reservations, location and variant; leave shortage visible or route to purchase."),
        ("Unit conversion missing", "The requested and native units have no approved conversion.", "Stop the transaction and ask the unit-master owner to add verified conversion evidence."),
        ("Item or variant not found", "The request and master identities do not match.", "Search code and variant; do not choose a similar name or create a duplicate casually."),
        ("Budget mapping unresolved", "Finance has not mapped the item or service to a budget head.", "Keep the commercial document in Draft and send the exact identity to Finance."),
        ("Budget unavailable", "The mapped head does not have enough available budget or the check failed.", "Do not bypass. Ask Budget or Finance to review allocation, period and commitment."),
        ("Receipt cannot be accepted", "PO, quantity, inspection, unit or location evidence is incomplete.", "Keep the receipt pending or quarantine the goods until the factual mismatch is resolved."),
        ("Service cannot be accepted", "Completion evidence is missing or the requester disagrees.", "Requester records rework required with a reason; Store does not create a GRN."),
    ],
    [1.65, 2.35, 3.05],
    8.6,
)

new_page()
doc.add_paragraph("Problem to action continued", style="Heading 1")
table(
    ["Problem or message", "What it usually means", "What to do next"],
    [
        ("Negative balance or stock exception", "A source movement, placement or document chain needs investigation.", "Open the exception, trace evidence, then correct by receipt, transfer, count, return or reversal."),
        ("Accepted stock has no location", "Receipt exists but put-away is incomplete.", "Use the receipt or put-away workflow; do not invent a warehouse on the movement."),
        ("Duplicate SKU or supplier", "A similar master identity already exists.", "Compare stable code, tax/legal identity, unit and lifecycle; merge or reactivate through the owner."),
        ("Valuation is partly unvalued", "Some stock lacks sufficient receipt or cost evidence.", "Expand affected items; repair receipt and landed-cost evidence. Unknown is not zero."),
        ("Totals differ after filtering", "The screen is showing a selected population, not the whole company or balance.", "Read the scope sentence and clear filters before reconciling totals."),
        ("Submit result is uncertain", "The response may have been interrupted after the server acted.", "Refresh and inspect status/history before resubmitting. Reuse retry only where the screen supports it."),
        ("Legacy record only", "The evidence predates the current workflow or lacks current company structure.", "Read it for history. Start all new demand in Request desk."),
        ("Page says retry but retry never helps", "The refusal needs identity, master data, permission or another owner.", "Capture the exact message and record number; send it to the owning team listed below."),
    ],
    [1.65, 2.35, 3.05],
    8.6,
)

new_page()
doc.add_paragraph("Status words that matter", style="Heading 1")
table(
    ["Area", "Status", "Employee meaning"],
    [
        ("Request", "Pending approval", "No fulfilment action should start yet."),
        ("Request", "Approved", "Store may classify and route lines according to authority."),
        ("Request", "Reserved / Picked / Issued", "Promised / physically staged / physically left Store - not interchangeable."),
        ("Purchase order", "Draft", "Editable preparation; not yet a supplier commitment."),
        ("Purchase order", "Issued", "Supplier-facing commitment; correct by supported amendment or cancellation."),
        ("Purchase order", "Partly received / Completed", "Some / all ordered physical quantities have valid receipt evidence."),
        ("Goods receipt", "Pending inspection", "Arrival recorded; quantity is not yet usable stock."),
        ("Goods receipt", "Accepted / Quarantined / Rejected", "Usable / isolated pending decision / not accepted into usable stock."),
        ("Service order", "Completion reported", "Supplier claims work is complete; requester must decide."),
        ("Service order", "Accepted / Rework required", "Requester accepts result / sends it back with a reason."),
        ("Stock count", "Open / Posted / Cancelled", "Counting in progress / approved differences recorded / no correction posted."),
        ("Valuation", "Partly unvalued / Indeterminate", "Some cost missing / evidence cannot support a reliable figure."),
    ],
    [1.25, 1.65, 4.15],
    8.7,
)

new_page()
doc.add_paragraph("Who owns the next action", style="Heading 1")
table(
    ["Issue", "Primary owner", "Send this evidence"],
    [
        ("Request purpose, quantity or approval", "Requester or approver", "Request number, line and required date"),
        ("Item, variant, category or base unit", "Material master owner", "Item code, proposed correction and source evidence"),
        ("Supplier identity or quotation", "Buyer or sourcing", "Supplier, offer date, currency, unit and validity"),
        ("Budget head or availability", "Budget or Finance", "Request/PO, item or service identity, amount and period"),
        ("Receipt quantity or condition", "Store receiver", "PO, receipt, line, physical count, photos where relevant"),
        ("Warehouse or location", "Inventory controller or administrator", "Warehouse, location, item and physical position"),
        ("Service result", "Requester", "Service order, completion evidence and rework reason"),
        ("Supplier voucher, tax or payment", "Accounting", "PO/service order, receipt/acceptance, invoice and charge details"),
        ("Role or staff-account link", "Administrator or HR", "Sign-in email, employee identity and required role"),
        ("System error", "Application support", "Exact message, page, record number, time and what happened before it"),
    ],
    [2.0, 1.8, 3.25],
    8.8,
)
doc.add_paragraph("Minimum escalation message", style="Heading 2")
p("State: I was on [page], working on [record number]. I selected [action]. The system showed [exact message] at [time]. I refreshed once and confirmed [current status]. No further submission was made.")
doc.add_paragraph("Use the full handbook when", style="Heading 2")
bullets([
    "You are onboarding to a role rather than completing one immediate task.",
    "A cross-application question involves Budget, Accounting, Production or costing policy.",
    "An unusual historical, audit, correction or reconciliation case is not covered by a task card.",
])


doc.core_properties.title = "GRAV Store Employee Quick Help"
doc.core_properties.subject = "Role based and task based Store and Purchase operating help"
doc.core_properties.author = "GRAV"
doc.core_properties.keywords = "Store, Purchase, Request, Receipt, Inventory, Service, Supplier, Budget, Troubleshooting"
doc.save(OUT)
print(OUT)
