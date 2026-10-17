#!/usr/bin/env python3
"""Build the complete speaker-aware Recording 5 Word transcript.

Input JSON is produced by the bounded diarization path.  This renderer keeps
ASR evidence intact and removes only exact word-boundary overlaps introduced by
adjacent decoding windows.
"""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


def clock(seconds: float) -> str:
    total = max(0, int(seconds))
    return f"{total // 3600:02d}:{(total % 3600) // 60:02d}:{total % 60:02d}"


def words(value: str) -> list[str]:
    return re.findall(r"\S+", value.strip())


def normal(word: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", word.lower())


def remove_exact_prefix_overlap(previous: str, current: str) -> str:
    """Remove only a verified adjacent ASR overlap, never rewrite speech."""
    earlier, later = words(previous), words(current)
    maximum = min(18, len(earlier), len(later))
    for count in range(maximum, 2, -1):
        if [normal(item) for item in earlier[-count:]] == [normal(item) for item in later[:count]]:
            return " ".join(later[count:]).strip()
    return current.strip()


def load_reviews(directory: Path) -> list[dict]:
    entries: list[dict] = []
    contextual_files = list(directory.glob("*-contextual-batch-*.json"))
    if contextual_files:
        for path in sorted(contextual_files):
            payload = json.loads(path.read_text(encoding="utf-8"))
            for window in payload.get("windows", []):
                text = str(window.get("romanText") or "").strip()
                if text:
                    entries.append({
                        "start": float(window["startSeconds"]),
                        "end": float(window["endSeconds"]),
                        "speaker": str(window.get("speakerActivity") or "Automatic transcript"),
                        "contextual": True,
                        "text": text,
                    })
        files = []
    else:
        files = list(directory.glob("recording-5-quality-chunk-*.json"))
    if not contextual_files and not files:
        files = list(directory.glob("recording-5-full-chunk-*.json"))
    for path in sorted(files, key=lambda item: int(re.search(r"(\d+)$", item.stem).group(1))):
        offset = float(re.search(r"(\d+)$", path.stem).group(1))
        payload = json.loads(path.read_text(encoding="utf-8"))
        for review in payload.get("utteranceReviews", []):
            text = str(review.get("selectedRomanText") or review.get("rawRomanText") or "").strip()
            if not text:
                continue
            entries.append({
                "start": offset + float(review["startSeconds"]),
                "end": offset + float(review["endSeconds"]),
                "speaker": str(review.get("speaker", "Speaker")),
                "text": text,
            })
    entries.sort(key=lambda item: (item["start"], item["end"], item["speaker"]))
    previous = ""
    for entry in entries:
        entry["text"] = remove_exact_prefix_overlap(previous, entry["text"])
        if entry["text"]:
            previous = entry["text"]
    return [entry for entry in entries if entry["text"]]


def add_run(paragraph, text: str, *, bold: bool = False, size: float | None = None, color=None):
    run = paragraph.add_run(text)
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    return run


def build(entries: list[dict], destination: Path, title_text: str) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    contextual = bool(entries) and all(entry.get("contextual") for entry in entries)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(title, "GRAV", bold=True, size=11, color=(0, 119, 91))
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(title, title_text, bold=True, size=22)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(
        subtitle,
        "Context-aware automatic transcript" if contextual else "Speaker-aware automatic transcript",
        size=10, color=(92, 99, 105),
    )

    note = document.add_paragraph()
    note.paragraph_format.space_after = Pt(14)
    add_run(note, "Timestamps refer to the original recording. ", bold=True, size=9, color=(92, 99, 105))
    add_run(
        note,
        "Each row is an overlapping audio-window decode; speaker labels show activity in the window, not word-level attribution."
        if contextual else "Speaker labels are assigned by the diarizer within each bounded recording segment.",
        size=9, color=(92, 99, 105),
    )

    active_block = None
    for entry in entries:
        block = int(entry["start"] // 180)
        if block != active_block:
            active_block = block
            heading = document.add_paragraph()
            heading.paragraph_format.space_before = Pt(12)
            heading.paragraph_format.space_after = Pt(5)
            add_run(heading, f"{clock(block * 180)} to {clock((block + 1) * 180)}", bold=True, size=12, color=(0, 119, 91))
        stamp = document.add_paragraph()
        stamp.paragraph_format.space_after = Pt(1)
        add_run(stamp, f"{clock(entry['start'])}  {entry['speaker']}", bold=True, size=9, color=(72, 80, 87))
        body = document.add_paragraph()
        body.paragraph_format.space_after = Pt(7)
        body.paragraph_format.line_spacing = 1.12
        add_run(body, entry["text"], size=10.5)

    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("input_directory", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--title", default="Recording 5 Transcript")
    args = parser.parse_args()
    entries = load_reviews(args.input_directory)
    if not entries:
        raise SystemExit("No transcript entries found")
    build(entries, args.output, args.title)
    print(f"Created {args.output} with {len(entries)} transcript entries")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
